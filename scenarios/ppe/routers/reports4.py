"""The four PPE reports (FRS-parity) + CSV/Excel export.

  1. Compliance     — per camera: checks, compliant, violations, compliance %.
  2. Violations     — each violation: time, camera, missing items, snapshot.
  3. By Item        — per PPE item (helmet/vest/..): violation count + last seen.
  4. Worker         — per worker track: compliant vs violations.

Each report has a JSON endpoint (UI table) and CSV / XLSX export via ?format=.
The snapshot column is embedded as an image in XLSX; CSV keeps a present/absent marker.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from pathlib import Path
from typing import Optional
from urllib.parse import parse_qs, urlparse

from fastapi import APIRouter, Depends, Query
from fastapi.responses import JSONResponse, Response
from sqlalchemy import and_, func, select

import config
from db import session
from db.models import PPEEvent
from deps import allowed_camera_ids, require_service_token
from routers._scope import apply_camera_scope
from schemas import iso, naive

router = APIRouter(tags=["reports"])

REPORTS = ("compliance", "violations", "by_item", "worker")
_VIOLATION_TYPES = ("ppe_missing", "ppe_removed")
_SAFE = re.compile(r"^[A-Za-z0-9\-_]+$")


# ── camera names (durable, FRS-parity) ─────────────────────────────────────
_CAM_CACHE: dict = {"at": 0.0, "map": {}}


def _camera_names() -> dict:
    """camera_id → name. Persisted to ppe_camera_names so names survive the core API
    going empty (scenario toggled off); refreshed from core when available (FRS parity)."""
    import time
    now = time.time()
    if now - _CAM_CACHE["at"] < 60 and _CAM_CACHE["map"]:
        return _CAM_CACHE["map"]
    names: dict = {}
    try:
        from db.models import PPECameraName
        with session() as s:
            for row in s.execute(select(PPECameraName)).scalars().all():
                names[str(row.camera_id)] = row.name
    except Exception:  # noqa: BLE001
        pass
    fresh: dict = {}
    try:
        from live.manager import _fetch_cameras
        for c in _fetch_cameras():
            cid = c.get("camera_id") or c.get("device_id") or c.get("id")
            nm = c.get("camera_name") or c.get("name")
            if cid and nm and str(cid) != str(nm):
                fresh[str(cid)] = nm
    except Exception:  # noqa: BLE001
        pass
    if fresh:
        names.update(fresh)
        try:
            from db.models import PPECameraName
            with session() as s:
                for cid, nm in fresh.items():
                    row = s.get(PPECameraName, cid)
                    if row is None:
                        s.add(PPECameraName(camera_id=cid, name=nm))
                    else:
                        row.name = nm
                s.commit()
        except Exception:  # noqa: BLE001
            pass
    if names:
        _CAM_CACHE.update(at=now, map=names)
    return names


def _cam(cid, names) -> str:
    if not cid:
        return "—"
    return names.get(str(cid)) or (str(cid)[:8])


# ── snapshot file resolver (for XLSX embed) ────────────────────────────────
def _snapshot_file(value: str) -> Optional[Path]:
    if not value:
        return None
    key = value
    if value.startswith("/snapshot") or value.startswith("http"):
        q = parse_qs(urlparse(value).query)
        key = (q.get("key") or [""])[0]
    frame_id = key.split(":", 1)[1] if ":" in key else key
    if not _SAFE.match(frame_id):
        return None
    base = config.DATA_PATH / "snapshots"
    for name in (f"{frame_id}_crop.jpg", f"{frame_id}.jpg"):
        p = base / name
        if p.exists():
            return p
    return None


# ── export helpers ─────────────────────────────────────────────────────────
def _csv(columns: list[str], rows: list[dict]) -> Response:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=columns, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        out = {c: r.get(c, "") for c in columns}
        if "snapshot" in columns:
            out["snapshot"] = "yes" if r.get("snapshot") else ""
        w.writerow(out)
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=report.csv"})


def _xlsx(columns: list[str], rows: list[dict], title: str = "Report") -> Response:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    has_snap = "snapshot" in columns
    ws.append([c.replace("_", " ").title() for c in columns])
    for c in ws[1]:
        c.font = Font(bold=True)
    THUMB = 56
    for ri, r in enumerate(rows, start=2):
        ws.append([("" if c == "snapshot" else r.get(c, "")) for c in columns])
        if not has_snap:
            continue
        src = _snapshot_file(r.get("snapshot") or "")
        if not src:
            continue
        try:
            img = XLImage(str(src))
            img.width = img.height = THUMB
            col = get_column_letter(columns.index("snapshot") + 1)
            ws.row_dimensions[ri].height = THUMB * 0.78
            ws.add_image(img, f"{col}{ri}")
        except Exception:  # noqa: BLE001
            pass
    if has_snap:
        ws.column_dimensions[get_column_letter(columns.index("snapshot") + 1)].width = 10
    bio = io.BytesIO()
    wb.save(bio)
    return Response(
        bio.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=report.xlsx"})


def _respond(columns, rows, fmt, title):
    fmt = (fmt or "json").lower()
    if fmt == "csv":
        return _csv(columns, rows)
    if fmt in ("xlsx", "excel"):
        return _xlsx(columns, rows, title)
    return JSONResponse({"columns": columns, "items": rows, "total": len(rows)})


def _range(day_from: str, day_to: str):
    start = naive(datetime.fromisoformat(day_from + "T00:00:00"))
    end = naive(datetime.fromisoformat(day_to + "T23:59:59"))
    return start, end


# ── 1. Compliance: per camera ──────────────────────────────────────────────
@router.get("/reports/compliance")
def report_compliance(day_from: str = Query(...), day_to: str = Query(...),
                      format: str = Query("json"),
                      _: None = Depends(require_service_token),
                      allowed: Optional[list[str]] = Depends(allowed_camera_ids)) -> Response:
    columns = ["camera", "checks", "compliant", "violations", "compliance_pct"]
    start, end = _range(day_from, day_to)
    names = _camera_names()
    with session() as s:
        conds = [PPEEvent.triggered_at >= start, PPEEvent.triggered_at <= end]
        if not apply_camera_scope(conds, PPEEvent.camera_id, None, allowed):
            return _respond(columns, [], format, "Compliance")
        rows_db = s.execute(
            select(PPEEvent.camera_id, PPEEvent.event_type, func.count())
            .where(and_(*conds)).group_by(PPEEvent.camera_id, PPEEvent.event_type)
        ).all()
        per: dict = {}
        for cid, et, n in rows_db:
            d = per.setdefault(cid, {"compliant": 0, "violations": 0})
            if et == "ppe_compliant":
                d["compliant"] += int(n)
            elif et in _VIOLATION_TYPES:
                d["violations"] += int(n)
        rows = []
        for cid, d in per.items():
            checks = d["compliant"] + d["violations"]
            pct = round(100.0 * d["compliant"] / checks, 1) if checks else 0.0
            rows.append({"camera": _cam(cid, names), "checks": checks,
                         "compliant": d["compliant"], "violations": d["violations"],
                         "compliance_pct": pct})
        rows.sort(key=lambda r: r["violations"], reverse=True)
    return _respond(columns, rows, format, "Compliance")


# ── 2. Violations: each missing/removed event ──────────────────────────────
@router.get("/reports/violations")
def report_violations(day_from: str = Query(...), day_to: str = Query(...),
                      format: str = Query("json"),
                      _: None = Depends(require_service_token),
                      allowed: Optional[list[str]] = Depends(allowed_camera_ids)) -> Response:
    columns = ["snapshot", "time", "camera", "missing", "type"]
    start, end = _range(day_from, day_to)
    names = _camera_names()
    with session() as s:
        conds = [PPEEvent.triggered_at >= start, PPEEvent.triggered_at <= end,
                 PPEEvent.event_type.in_(_VIOLATION_TYPES)]
        if not apply_camera_scope(conds, PPEEvent.camera_id, None, allowed):
            return _respond(columns, [], format, "Violations")
        evs = s.execute(select(PPEEvent).where(and_(*conds))
                        .order_by(PPEEvent.triggered_at.desc()).limit(2000)).scalars().all()
        rows = []
        for e in evs:
            miss = e.missing_items or ([e.ppe_item] if e.ppe_item else [])
            rows.append({
                "snapshot": e.snapshot_path or "",
                "time": iso(e.triggered_at),
                "camera": _cam(e.camera_id, names),
                "missing": ", ".join(str(m) for m in miss) or "—",
                "type": "Removed" if e.event_type == "ppe_removed" else "Missing",
            })
    if (format or "json").lower() == "json":
        return JSONResponse({"columns": columns, "items": rows, "total": len(rows)})
    return _respond(columns, rows, format, "Violations")


# ── 3. By Item: per PPE item ───────────────────────────────────────────────
@router.get("/reports/by_item")
def report_by_item(day_from: str = Query(...), day_to: str = Query(...),
                   format: str = Query("json"),
                   _: None = Depends(require_service_token),
                   allowed: Optional[list[str]] = Depends(allowed_camera_ids)) -> Response:
    columns = ["item", "violations", "last_seen"]
    start, end = _range(day_from, day_to)
    with session() as s:
        conds = [PPEEvent.triggered_at >= start, PPEEvent.triggered_at <= end,
                 PPEEvent.event_type.in_(_VIOLATION_TYPES)]
        if not apply_camera_scope(conds, PPEEvent.camera_id, None, allowed):
            return _respond(columns, [], format, "By Item")
        evs = s.execute(select(PPEEvent.missing_items, PPEEvent.ppe_item,
                               PPEEvent.triggered_at)
                        .where(and_(*conds))).all()
        agg: dict = {}
        for miss, item, ts in evs:
            items = miss or ([item] if item else [])
            for it in items:
                d = agg.setdefault(str(it).lower(), {"n": 0, "last": None})
                d["n"] += 1
                if d["last"] is None or (ts and ts > d["last"]):
                    d["last"] = ts
        rows = [{"item": k.title(), "violations": v["n"], "last_seen": iso(v["last"])}
                for k, v in sorted(agg.items(), key=lambda kv: kv[1]["n"], reverse=True)]
    return _respond(columns, rows, format, "By Item")


# ── 4. Worker: per worker track ────────────────────────────────────────────
@router.get("/reports/worker")
def report_worker(day_from: str = Query(...), day_to: str = Query(...),
                  format: str = Query("json"),
                  _: None = Depends(require_service_token),
                  allowed: Optional[list[str]] = Depends(allowed_camera_ids)) -> Response:
    columns = ["worker", "violations", "compliant", "last_seen"]
    start, end = _range(day_from, day_to)
    with session() as s:
        conds = [PPEEvent.triggered_at >= start, PPEEvent.triggered_at <= end,
                 PPEEvent.worker_track_id.isnot(None)]
        if not apply_camera_scope(conds, PPEEvent.camera_id, None, allowed):
            return _respond(columns, [], format, "Worker")
        rows_db = s.execute(
            select(PPEEvent.worker_track_id, PPEEvent.event_type,
                   func.count(), func.max(PPEEvent.triggered_at))
            .where(and_(*conds))
            .group_by(PPEEvent.worker_track_id, PPEEvent.event_type)
        ).all()
        per: dict = {}
        for wid, et, n, last in rows_db:
            d = per.setdefault(wid, {"violations": 0, "compliant": 0, "last": None})
            if et == "ppe_compliant":
                d["compliant"] += int(n)
            elif et in _VIOLATION_TYPES:
                d["violations"] += int(n)
            if d["last"] is None or (last and last > d["last"]):
                d["last"] = last
        rows = [{"worker": f"Worker #{wid}", "violations": d["violations"],
                 "compliant": d["compliant"], "last_seen": iso(d["last"])}
                for wid, d in sorted(per.items(), key=lambda kv: kv[1]["violations"], reverse=True)]
    return _respond(columns, rows, format, "Worker")
